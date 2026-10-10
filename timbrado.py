"""Papel timbrado e helpers python-docx compartilhados pelos documentos da
ONG Santidade Integral em Cristo.

Uso nos geradores (que vivem em subpastas):

    import sys; from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from timbrado import *
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
LOGO = ROOT / "assets" / "logo.png"
ONG = "SANTIDADE INTEGRAL EM CRISTO"
ONG_SUB = "Ação Social  •  Assistência e Promoção Familiar"

GREEN = RGBColor(0x2D, 0x51, 0x30)
GREEN_HEX = "2D5130"
LIGHT_HEX = "E8EFE8"
GREY = RGBColor(0x59, 0x59, 0x59)
FONT = "Calibri"
BOX = "☐"  # ☐


# ----------------------------------------------------------------- helpers
def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def borders(table, color="A6B8A6", size=4):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        b.append(el)
    tbl_pr.append(b)


def para_bottom_border(paragraph, color=GREEN_HEX, size=12):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    p_pr.append(bdr)


def run(p, text, bold=False, italic=False, size=9.5, color=None, font=FONT):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color is not None:
        r.font.color.rgb = color
    return r


def spacing(p, before=0, after=2, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def section_title(doc, number, title):
    p = doc.add_paragraph()
    spacing(p, before=8, after=3)
    p.paragraph_format.keep_with_next = True
    para_bottom_border(p)
    run(p, f"{number}  ", bold=True, size=11, color=GREEN)
    run(p, title.upper(), bold=True, size=11, color=GREEN)
    return p


LINE_CAP = 104  # capacidade aproximada da linha em caracteres (Calibri 9.5pt, 17.4cm)


def field_line(doc, *items, size=9.5):
    """items: ("Rótulo", peso) → rótulo + linha de preenchimento proporcional; ou texto puro."""
    p = doc.add_paragraph()
    spacing(p, after=3)
    labels = [it for it in items if isinstance(it, tuple)]
    fixed = sum(len(it[0]) + 2 for it in labels) + sum(len(it) for it in items if not isinstance(it, tuple))
    fixed += 3 * (len(items) - 1)
    free = max(LINE_CAP - fixed, 6 * len(labels))
    total_w = sum(it[1] for it in labels) or 1
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            label, w = it
            run(p, f"{label}: ", bold=True, size=size)
            run(p, "_" * max(4, int(free * w / total_w)), size=size, color=GREY)
        else:
            run(p, it, size=size)
        if i < len(items) - 1:
            run(p, "   ", size=size)
    return p


def check_line(doc, label, options, size=9, note=None):
    p = doc.add_paragraph()
    spacing(p, after=3)
    run(p, f"{label}: ", bold=True, size=size)
    for o in options:
        run(p, f"{BOX} {o}  ", size=size)
    if note:
        run(p, note, size=size, color=GREY)
    return p


def table(doc, headers, rows, widths_cm, header_fill=GREEN_HEX, size=8.5, blank_rows=0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    set_cell_margins(t)
    for i, (h, w) in enumerate(zip(headers, widths_cm)):
        c = t.rows[0].cells[i]
        c.width = Cm(w)
        shade(c, header_fill)
        p = c.paragraphs[0]
        spacing(p, after=0)
        run(p, h, bold=True, size=size, color=RGBColor(0xFF, 0xFF, 0xFF))
    all_rows = list(rows) + [[""] * len(headers)] * blank_rows
    for r in all_rows:
        cells = t.add_row().cells
        for i, (val, w) in enumerate(zip(r, widths_cm)):
            cells[i].width = Cm(w)
            p = cells[i].paragraphs[0]
            spacing(p, after=0)
            if isinstance(val, tuple):  # (text, bold)
                run(p, val[0], bold=val[1], size=size)
            else:
                run(p, val, size=size)
    # altura mínima para linhas em branco
    for row in t.rows[1:]:
        tr_pr = row._tr.get_or_add_trPr()
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), "300")
        h.set(qn("w:hRule"), "atLeast")
        tr_pr.append(h)
    fix_widths(t, widths_cm)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def note(doc, text, size=8.5):
    p = doc.add_paragraph()
    spacing(p, after=3)
    run(p, text, italic=True, size=size, color=GREY)
    return p


def lines_box(doc, n_lines=3, label=None):
    if label:
        p = doc.add_paragraph()
        spacing(p, after=2)
        run(p, label, bold=True, size=9.5)
    for _ in range(n_lines):
        p = doc.add_paragraph()
        spacing(p, before=9, after=0, line=1.0)
        para_bottom_border(p, color="9A9A9A", size=4)
        run(p, " ", size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def fix_widths(t, widths_cm):
    """LibreOffice só respeita larguras se o grid e todas as células forem definidos."""
    tbl_pr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tbl_pr.append(lay)
    for col, w in zip(t.columns, widths_cm):
        col.width = Cm(w)
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)



def add_page_field(paragraph, size=7.5, color=GREY):
    r = paragraph.add_run()
    r.font.size = Pt(size)
    r.font.color.rgb = color
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)


def new_document(*, header_meta="", footer_label="", top_margin_cm=3.9, bottom_margin_cm=1.6):
    """Documento A4 com o papel timbrado da ONG (logo no canto superior esquerdo).

    header_meta: linha cinza opcional abaixo do nome da ONG (campos do cabeçalho).
    footer_label: texto do rodapé antes de "Página N".
    """
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = Cm(top_margin_cm)
    sec.bottom_margin = Cm(bottom_margin_cm)
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.6)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(9.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    header = sec.header
    header.is_linked_to_previous = False
    ht = header.add_table(rows=1, cols=2, width=Cm(17.4))
    ht.autofit = False
    fix_widths(ht, [3.1, 14.3])
    tr_pr = ht.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(2.6 * 567)))
    h.set(qn("w:hRule"), "atLeast")
    tr_pr.append(h)
    hp = ht.rows[0].cells[0].paragraphs[0]
    spacing(hp, after=0)
    hp.add_run().add_picture(str(LOGO), height=Cm(2.4))
    c1 = ht.rows[0].cells[1]
    p = c1.paragraphs[0]
    spacing(p, before=4, after=0)
    run(p, ONG, bold=True, size=11, color=GREEN)
    p = c1.add_paragraph()
    spacing(p, after=0)
    run(p, ONG_SUB, size=8.5, color=GREY)
    if header_meta:
        p = c1.add_paragraph()
        spacing(p, after=0)
        run(p, header_meta, size=8.5, color=GREY)
    hl = header.add_paragraph()
    spacing(hl, after=0)
    para_bottom_border(hl, size=8)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(fp, after=0)
    run(fp, f"{footer_label}  •  Página ", size=7.5, color=GREY)
    add_page_field(fp)
    return doc


def title_block(doc, title, subtitle=None, verse=None, verse_ref=None, after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, after=0)
    run(p, title, bold=True, size=16, color=GREEN)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, after=0)
        run(p, subtitle, size=10, color=GREY)
    if verse:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, before=4, after=after)
        # "\n" no versículo = quebra de linha manual (pra equilibrar versos longos)
        parts = verse.split("\n")
        for i, part in enumerate(parts):
            txt = ("“" if i == 0 else "") + part + ("”  " if i == len(parts) - 1 else "")
            r = run(p, txt, italic=True, size=10, color=GREEN)
            if i < len(parts) - 1:
                r.add_break()
        run(p, verse_ref, bold=True, size=10, color=GREEN)


def body(doc, text, size=9.5, after=4, justify=True, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    spacing(p, after=after)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(p, text, size=size, bold=bold, italic=italic, color=color)
    return p


def bullet(doc, text, size=9.5, after=2, lead=None):
    """Item com marcador simples (sem depender de estilos de lista do Word)."""
    p = doc.add_paragraph()
    spacing(p, after=after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    run(p, "•  ", size=size, color=GREEN, bold=True)
    if lead:
        run(p, lead + " ", size=size, bold=True)
    run(p, text, size=size)
    return p


def signature_row(doc, labels, widths_cm=None, before=18):
    n = len(labels)
    widths_cm = widths_cm or [17.4 / n] * n
    t = doc.add_table(rows=1, cols=n)
    t.autofit = False
    fix_widths(t, widths_cm)
    for i, label in enumerate(labels):
        c = t.rows[0].cells[i]
        p = c.paragraphs[0]
        spacing(p, before=before, after=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, "_" * int(widths_cm[i] * 4.6), color=GREY)
        p = c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, after=0)
        run(p, label, size=8.5, color=GREY)
    return t


def to_pdf(docx_path):
    """Converte via LibreOffice headless; devolve o caminho do PDF."""
    import subprocess
    docx_path = Path(docx_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                    str(docx_path.parent), str(docx_path)],
                   check=True, capture_output=True, timeout=180)
    return docx_path.with_suffix(".pdf")
